"""人工決策表 (decisions.xlsx)。

這是整個 pipeline 唯一允許人手編輯的檔案，也是同仁參與的入口。
統整表本身是產出物，每次重跑都會覆蓋；任何人工判斷都必須寫在這裡才會留存。

每張決策表都要求填寫「覆核者」與「覆核日期」，理由欄則會原樣帶進統整表的
備註，讓報告的每一個人工介入都可追溯到人與時間 —— 這是 GLP 稽核會問的東西。
"""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date
from pathlib import Path
from typing import Any

import pandas as pd

SHEET_FINAL = "最終採用覆核"
SHEET_HIGHSD = "HIGHSD單孔採用"
SHEET_GROUP = "動物分組指定"
SHEET_ORGAN = "臟器代碼補充"
SHEET_README = "填寫說明"

KEY_COLUMNS = ["動物編號", "臟器代碼", "來源檔案"]
AUDIT_COLUMNS = ["覆核者", "覆核日期", "理由"]

SCHEMA: dict[str, list[str]] = {
    SHEET_FINAL: KEY_COLUMNS + ["最終採用(Y/N)"] + AUDIT_COLUMNS,
    SHEET_HIGHSD: KEY_COLUMNS + ["採用方式(孔位1/孔位2/兩孔平均)"] + AUDIT_COLUMNS,
    SHEET_GROUP: ["動物編號", "組別"] + AUDIT_COLUMNS,
    SHEET_ORGAN: ["臟器代碼", "臟器名稱(英文)", "臟器名稱(中文)"] + AUDIT_COLUMNS,
}

HIGHSD_CHOICES = {"孔位1", "孔位2", "兩孔平均"}

README_LINES = [
    ["decisions.xlsx — 人工決策表"],
    [""],
    ["這是整個統整流程唯一需要人工填寫的檔案。統整表(輸出)請勿手改，重跑會被覆蓋。"],
    [""],
    ["分頁用途："],
    [f"  {SHEET_FINAL}：同一動物＋臟器有多個版本時，指定哪一列為最終採用。"],
    ["      不填 = 依系統預設規則（rerun 優先於原始）。"],
    [f"  {SHEET_HIGHSD}：兩重複孔位差異過大(HIGHSD=Y)且無 rerun 時，指定採用哪一孔。"],
    ["      不填 = 系統不自行選孔，該列標記為「請人工複核」並以兩孔平均暫呈。"],
    [f"  {SHEET_GROUP}：時間點無法唯一判定組別時（例如 Day 29 兩組皆可能採樣），"],
    ["      在此指定該動物的組別。"],
    [f"  {SHEET_ORGAN}：新出現、尚未收錄於設定檔的臟器代碼，在此補上名稱。"],
    [""],
    ["填寫規則："],
    ["  1. 動物編號與臟器代碼請保留前導 0（例如 0006、03），欄位已設為文字格式。"],
    ["  2. 覆核者與覆核日期為必填；缺一者該列會被忽略並在執行時提示，不會靜默套用。"],
    ["  3. 理由會原樣帶進統整表備註欄，供追溯與稽核。"],
    ["  4. 只填你要覆寫的列，其餘留空即可。"],
]


@dataclass
class DecisionSet:
    """讀入並驗證過的決策集合。"""

    final_use: dict[tuple[str, str, str], dict[str, Any]] = field(default_factory=dict)
    highsd_well: dict[tuple[str, str, str], dict[str, Any]] = field(default_factory=dict)
    group_override: dict[str, dict[str, Any]] = field(default_factory=dict)
    organ_override: dict[str, dict[str, Any]] = field(default_factory=dict)
    warnings: list[str] = field(default_factory=list)

    @property
    def count(self) -> int:
        return (
            len(self.final_use) + len(self.highsd_well)
            + len(self.group_override) + len(self.organ_override)
        )


def _norm(value: Any) -> str:
    """統一鍵值格式。Excel 常把 '0006' 讀成數字 6，這裡把它還原。"""
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _norm_id(value: Any, width: int) -> str:
    """還原前導 0 的識別碼（動物編號 4 碼、臟器代碼 2 碼）。"""
    text = _norm(value)
    if text and text.isdigit() and len(text) < width:
        return text.zfill(width)
    return text


def _audit_ok(row: pd.Series, sheet: str, label: str,
              warnings: list[str]) -> tuple[bool, dict[str, Any]]:
    reviewer = _norm(row.get("覆核者"))
    reviewed_on = row.get("覆核日期")
    reason = _norm(row.get("理由"))

    if not reviewer or _norm(reviewed_on) == "":
        warnings.append(
            f"[{sheet}] {label}：缺少覆核者或覆核日期，此列已忽略（不會套用）。"
        )
        return False, {}

    if isinstance(reviewed_on, (pd.Timestamp, date)):
        reviewed_text = pd.Timestamp(reviewed_on).date().isoformat()
    else:
        reviewed_text = _norm(reviewed_on)

    return True, {"覆核者": reviewer, "覆核日期": reviewed_text, "理由": reason}


def create_template(path: str | Path, *, overwrite: bool = False) -> Path:
    """建立空白決策表。已存在時預設不覆蓋（避免洗掉同仁填的內容）。"""
    target = Path(path)
    if target.exists() and not overwrite:
        return target
    target.parent.mkdir(parents=True, exist_ok=True)

    with pd.ExcelWriter(target, engine="openpyxl") as writer:
        pd.DataFrame(README_LINES).to_excel(
            writer, sheet_name=SHEET_README, index=False, header=False
        )
        for sheet, columns in SCHEMA.items():
            pd.DataFrame(columns=columns).to_excel(writer, sheet_name=sheet, index=False)

    _format_template(target)
    return target


def _format_template(path: Path) -> None:
    """把識別碼欄位設成文字格式，否則 Excel 會吃掉 '0006' 的前導 0。"""
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    workbook = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="DDEBF7")
    header_font = Font(bold=True)

    for sheet_name, columns in SCHEMA.items():
        sheet = workbook[sheet_name]
        for index, column in enumerate(columns, start=1):
            cell = sheet.cell(row=1, column=index)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(vertical="center")
            sheet.column_dimensions[cell.column_letter].width = max(14, len(column) + 4)
            if column in {"動物編號", "臟器代碼"}:
                for row in range(2, 500):
                    sheet.cell(row=row, column=index).number_format = "@"
        sheet.freeze_panes = "A2"

    readme = workbook[SHEET_README]
    readme.column_dimensions["A"].width = 90
    readme["A1"].font = Font(bold=True, size=13)
    workbook.save(path)


def load_decisions(path: str | Path) -> DecisionSet:
    """讀取決策表。檔案不存在時回傳空集合（不是錯誤 —— 第一次跑本來就沒有）。"""
    decisions = DecisionSet()
    target = Path(path)
    if not target.is_file():
        decisions.warnings.append(
            f"找不到決策表 {target}，本次以系統預設規則執行。"
            "（可用 `qpcr-biod init-decisions` 產生空白範本）"
        )
        return decisions

    sheets = pd.read_excel(target, sheet_name=None, dtype=object)

    for sheet_name in SCHEMA:
        if sheet_name not in sheets:
            decisions.warnings.append(f"決策表缺少分頁「{sheet_name}」，該類決策視為未填。")

    _load_final_use(sheets.get(SHEET_FINAL), decisions)
    _load_highsd(sheets.get(SHEET_HIGHSD), decisions)
    _load_group(sheets.get(SHEET_GROUP), decisions)
    _load_organ(sheets.get(SHEET_ORGAN), decisions)
    return decisions


def _iter_rows(frame: pd.DataFrame | None):
    if frame is None or frame.empty:
        return
    for _, row in frame.iterrows():
        if row.isna().all():
            continue
        yield row


def _load_final_use(frame: pd.DataFrame | None, decisions: DecisionSet) -> None:
    for row in _iter_rows(frame):
        animal = _norm_id(row.get("動物編號"), 4)
        organ = _norm_id(row.get("臟器代碼"), 2)
        source = _norm(row.get("來源檔案"))
        choice = _norm(row.get("最終採用(Y/N)")).upper()
        label = f"{animal}_{organ} @ {source or '(未指定檔案)'}"

        if not (animal and organ and source):
            decisions.warnings.append(f"[{SHEET_FINAL}] 動物編號/臟器代碼/來源檔案需三者齊全，此列已忽略：{label}")
            continue
        if choice not in {"Y", "N"}:
            decisions.warnings.append(f"[{SHEET_FINAL}] {label}：「最終採用」需填 Y 或 N，此列已忽略。")
            continue
        ok, audit = _audit_ok(row, SHEET_FINAL, label, decisions.warnings)
        if not ok:
            continue
        decisions.final_use[(animal, organ, source)] = {"最終採用": choice, **audit}


def _load_highsd(frame: pd.DataFrame | None, decisions: DecisionSet) -> None:
    for row in _iter_rows(frame):
        animal = _norm_id(row.get("動物編號"), 4)
        organ = _norm_id(row.get("臟器代碼"), 2)
        source = _norm(row.get("來源檔案"))
        choice = _norm(row.get("採用方式(孔位1/孔位2/兩孔平均)"))
        label = f"{animal}_{organ} @ {source or '(未指定檔案)'}"

        if not (animal and organ and source):
            decisions.warnings.append(f"[{SHEET_HIGHSD}] 動物編號/臟器代碼/來源檔案需三者齊全，此列已忽略：{label}")
            continue
        if choice not in HIGHSD_CHOICES:
            decisions.warnings.append(
                f"[{SHEET_HIGHSD}] {label}：採用方式需為 {'/'.join(sorted(HIGHSD_CHOICES))} 之一，此列已忽略。"
            )
            continue
        ok, audit = _audit_ok(row, SHEET_HIGHSD, label, decisions.warnings)
        if not ok:
            continue
        decisions.highsd_well[(animal, organ, source)] = {"採用方式": choice, **audit}


def _load_group(frame: pd.DataFrame | None, decisions: DecisionSet) -> None:
    for row in _iter_rows(frame):
        animal = _norm_id(row.get("動物編號"), 4)
        group = _norm(row.get("組別"))
        if not animal or not group:
            continue
        ok, audit = _audit_ok(row, SHEET_GROUP, animal, decisions.warnings)
        if not ok:
            continue
        decisions.group_override[animal] = {"組別": group, **audit}


def _load_organ(frame: pd.DataFrame | None, decisions: DecisionSet) -> None:
    for row in _iter_rows(frame):
        code = _norm_id(row.get("臟器代碼"), 2)
        if not code:
            continue
        ok, audit = _audit_ok(row, SHEET_ORGAN, code, decisions.warnings)
        if not ok:
            continue
        decisions.organ_override[code] = {
            "en": _norm(row.get("臟器名稱(英文)")),
            "zh": _norm(row.get("臟器名稱(中文)")),
            **audit,
        }
