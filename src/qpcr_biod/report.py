"""輸出統整活頁簿。"""

from __future__ import annotations

from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .config import StudyConfig
from .consolidate import (
    STYLE_EXPECTED_ND,
    STYLE_MANUAL_REVIEW,
    STYLE_NOT_FINAL,
)
from .lob import LOBResult
from .manifest import Manifest
from .summary import SummaryTables

ROW_FILLS = {
    STYLE_MANUAL_REVIEW: PatternFill("solid", fgColor="FCE4E4"),  # 淺紅：請人工複核
    STYLE_EXPECTED_ND: PatternFill("solid", fgColor="E2EFDA"),    # 淺綠：ND，屬預期現象
    STYLE_NOT_FINAL: PatternFill("solid", fgColor="F2F2F2"),      # 淺灰：非最終採用
}

HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")
HEADER_FONT = Font(bold=True)

SHEET_NOTES = "說明"
SHEET_MAIN = "統整數據"
SHEET_ORGAN_SUMMARY = "Organ Mean SD Summary"
SHEET_BLOOD_LOB = "血液LOB校正"
SHEET_ORGAN_CODES = "臟器代碼對照表"
SHEET_CURVES = "標準曲線與QC"
SHEET_ANIMALS = "動物分組與性別對照表"
SHEET_DOTPLOT = "Dot_Plot_Data"
SHEET_RAW = "Raw_Import"
SHEET_FILES = "已處理檔案紀錄"
SHEET_MANIFEST = "執行紀錄"

# 這些欄位是內部用的，不輸出到報表
INTERNAL_COLUMNS = ["列標記"]


def write_report(path: str | Path, *, config: StudyConfig,
                 consolidated: pd.DataFrame, summaries: SummaryTables,
                 lob: LOBResult, curves_frame: pd.DataFrame,
                 qc_frame: pd.DataFrame, organ_codes: pd.DataFrame,
                 animals: pd.DataFrame, raw: pd.DataFrame,
                 files: pd.DataFrame, manifest: Manifest) -> Path:
    """把所有分頁寫成一個活頁簿。"""
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)

    display = consolidated.drop(columns=INTERNAL_COLUMNS, errors="ignore")

    with pd.ExcelWriter(target, engine="openpyxl") as writer:
        _write(writer, SHEET_NOTES, _notes_frame(config, manifest, summaries, lob),
               header=False)
        _write(writer, SHEET_MAIN, display)
        _write(writer, SHEET_ORGAN_SUMMARY, summaries.organ_summary)

        for name, frame in summaries.subsets.items():
            _write(writer, _safe_sheet_name(name), frame)

        if lob.enabled:
            _write(writer, SHEET_BLOOD_LOB, _lob_frame(lob, summaries))

        _write(writer, SHEET_CURVES, curves_frame)
        _write(writer, "QC點位明細", qc_frame)
        _write(writer, SHEET_ORGAN_CODES, organ_codes)
        _write(writer, SHEET_ANIMALS, animals)
        if not summaries.dot_plot.empty:
            _write(writer, SHEET_DOTPLOT, summaries.dot_plot)
        _write(writer, SHEET_FILES, files)
        _write(writer, SHEET_MANIFEST, manifest.to_frame(), header=False)
        _write(writer, "執行紀錄-輸入檔", manifest.inputs_frame())
        _write(writer, "執行紀錄-人工決策", manifest.decisions_frame())
        _write(writer, "執行紀錄-警告", manifest.warnings_frame())
        _write(writer, SHEET_RAW, raw)

    _apply_formatting(target, consolidated)
    return target


def _write(writer: pd.ExcelWriter, sheet: str, frame: pd.DataFrame,
           header: bool = True) -> None:
    if frame is None:
        frame = pd.DataFrame()
    if frame.empty and header:
        frame = pd.DataFrame({"（本次無資料）": []})
    frame.to_excel(writer, sheet_name=sheet, index=False, header=header)


def _safe_sheet_name(name: str) -> str:
    """Excel 分頁名不得含 : \\ / ? * [ ]，且上限 31 字元。"""
    cleaned = "".join("-" if ch in ':\\/?*[]' else ch for ch in str(name))
    return cleaned[:31] or "Sheet"


def _notes_frame(config: StudyConfig, manifest: Manifest,
                 summaries: SummaryTables, lob: LOBResult) -> pd.DataFrame:
    eff_low, eff_high = config.standard_curve.get("efficiency_range", [85.0, 115.0])
    rec_low, rec_high = config.qc.get("sensitivity_control", {}).get(
        "recovery_range", [80.0, 120.0]
    )
    lines: list[str] = [
        f"{config.study_title}",
        "",
        f"研究編號：{config.study_id}",
        f"產生時間：{manifest.generated_at}（執行者 {manifest.generated_by}）",
        f"程式版本：{manifest.tool_version}",
        f"單位：{config.unit}",
        "",
        "【這份檔案是產出物，請勿手動修改】",
        "本活頁簿由 qpcr-biod pipeline 從原始 StepOnePlus export 重新計算產生，",
        "每次重跑都會整份覆蓋。任何人工判斷（rerun 採用、HIGHSD 選孔、組別指定、",
        "臟器代碼補充）請填在 decisions.xlsx，重跑後會自動套用並記錄於「執行紀錄」。",
        "",
        "分頁說明：",
        f"  {SHEET_MAIN}：每列 = 一個動物編號＋臟器代碼＋來源檔案（含原始/rerun 版本）。",
        "      「呈現結果」為建議之報告呈現值：低於該 run LLOQ 者標示 ND。",
        "      「原始測得平均值(未覆蓋)」保留儀器測得值，供追溯查核。",
        "      底色：淺紅=請人工複核；淺綠=ND 屬預期現象；淺灰=非最終採用（保留供追溯）。",
        f"  {SHEET_ORGAN_SUMMARY}：以臟器為主體、男女分列、依時間點排序之 Mean±SD 報告表。",
        "      Mean/SD 僅納入可定量(非 ND)數值；ND 檢體數記於 Remarks，不予估算代入。",
        f"  {SHEET_CURVES} / QC點位明細：每個 run 的回歸統計與允收判定。",
        f"      允收標準：R² ≥ {config.standard_curve.get('min_r_squared', 0.98)}、"
        f"PCR 效率 {eff_low}–{eff_high}%。",
        f"      回收率以敏感度對照組為唯一判定依據（{rec_low}–{rec_high}%）。",
        f"  {SHEET_RAW}：所有原始檔完整匯入之逐孔資料，為本活頁簿所有數值的來源。",
        f"  {SHEET_FILES}：已匯入的原始檔案清單。",
        "  執行紀錄／-輸入檔／-人工決策／-警告：本次執行的完整追溯資訊，",
        "      含每個輸入檔的 SHA-256，可用於驗證產出是否可重現。",
    ]

    if lob.enabled:
        lines += ["", f"  {SHEET_BLOOD_LOB}：血液(臟器代碼 {lob.organ_code}) 的 LOB 背景校正。"]
        if lob.usable:
            lines += [
                f"      LOB 由 {lob.baseline_timepoint} 血液孔位層級資料建立"
                f"（n={lob.n_wells}，Mean+{lob.primary_multiplier}×SD = {lob.lob_primary:.6f} pg）。",
                "      陽性需同時通過『非 ND』與『≥ LOB』；"
                f"{lob.baseline_timepoint} 為 LOB 定義組，其欄位標示為 - 不予計算。",
                "      本頁為排查/篩選用途，不變更統整數據既有 ND 判定或任何原始數值；",
                "      正式採用前仍須與研究人員/QA 確認。",
            ]
        else:
            lines += [f"      本次未建立 LOB：{'；'.join(lob.notes) or '條件不足'}"]

    if summaries.notes:
        lines += ["", "其他註記："] + [f"  • {note}" for note in summaries.notes]

    if manifest.warnings:
        lines += ["", f"本次執行有 {len(manifest.warnings)} 則警告，詳見「執行紀錄-警告」分頁。"]

    return pd.DataFrame(lines)


def _lob_frame(lob: LOBResult, summaries: SummaryTables) -> pd.DataFrame:
    """把 LOB 的計算依據、比對明細與校正後彙整疊成一張分頁。"""
    blocks: list[pd.DataFrame] = []

    header = pd.DataFrame([
        ["A. LOB 計算依據（Predose 血液孔位層級資料）"],
        [f"n（孔位數）= {lob.n_wells}"],
        [f"Mean (pg) = {lob.mean:.6f}" if lob.mean is not None else "Mean = 無法計算"],
        [f"SD (pg) = {lob.sd:.6f}" if lob.sd is not None else "SD = 無法計算"],
        [f"LOB = Mean + {lob.primary_multiplier}×SD = "
         f"{lob.lob_primary:.6f} pg（本次分析採用值）"
         if lob.lob_primary is not None else "LOB 未建立"],
        [f"LOB = Mean + {lob.reference_multiplier}×SD = "
         f"{lob.lob_reference:.6f} pg（僅供參考）"
         if lob.lob_reference is not None else ""],
    ])
    blocks.append(header)
    for note in lob.notes:
        blocks.append(pd.DataFrame([[f"※ {note}"]]))

    blocks.append(pd.DataFrame([[""], ["B. Predose 血液孔位明細"]]))
    blocks.append(_frame_with_header(lob.wells_detail))

    blocks.append(pd.DataFrame([[""], ["C. 各血液檢體與 LOB 比對（僅列最終採用列）"]]))
    blocks.append(_frame_with_header(lob.comparison))

    if not summaries.blood_lob_adjusted.empty:
        blocks.append(pd.DataFrame([[""], ["D. LOB-Adjusted Calculation（校正後彙整）"]]))
        blocks.append(_frame_with_header(summaries.blood_lob_adjusted))

    return pd.concat(blocks, ignore_index=True)


def _frame_with_header(frame: pd.DataFrame) -> pd.DataFrame:
    """把 DataFrame 攤平成含表頭列的無名欄位表，方便多塊疊在同一分頁。"""
    if frame is None or frame.empty:
        return pd.DataFrame([["（無資料）"]])
    rows = [list(frame.columns)]
    for _, row in frame.iterrows():
        rows.append(list(row.values))
    return pd.DataFrame(rows)


def _apply_formatting(path: Path, consolidated: pd.DataFrame) -> None:
    from openpyxl import load_workbook

    workbook = load_workbook(path)

    for sheet in workbook.worksheets:
        if sheet.max_row < 1:
            continue
        if sheet.title not in {SHEET_NOTES, SHEET_MANIFEST}:
            for cell in sheet[1]:
                if cell.value is not None:
                    cell.fill = HEADER_FILL
                    cell.font = HEADER_FONT
                    cell.alignment = Alignment(vertical="center", wrap_text=True)
            sheet.freeze_panes = "A2"
        _autosize(sheet)

    if SHEET_MAIN in workbook.sheetnames and "列標記" in consolidated.columns:
        sheet = workbook[SHEET_MAIN]
        width = sheet.max_column
        for offset, style in enumerate(consolidated["列標記"].tolist(), start=2):
            fill = ROW_FILLS.get(style)
            if fill is None:
                continue
            for column in range(1, width + 1):
                sheet.cell(row=offset, column=column).fill = fill

    _force_text_identifiers(workbook)

    if SHEET_NOTES in workbook.sheetnames:
        notes = workbook[SHEET_NOTES]
        notes.column_dimensions["A"].width = 100
        notes["A1"].font = Font(bold=True, size=13)

    workbook.save(path)


# 這些欄位是識別碼不是數字，"03" 的前導 0 必須留著
IDENTIFIER_HEADERS = {
    "動物編號", "臟器代碼", "代碼", "Organ code", "Animal No", "孔位",
    "孔位1-Well", "孔位2-Well", "Well",
}


def _force_text_identifiers(workbook) -> None:
    for sheet in workbook.worksheets:
        if sheet.max_row < 2:
            continue
        for cell in sheet[1]:
            if cell.value in IDENTIFIER_HEADERS:
                for row in range(2, sheet.max_row + 1):
                    sheet.cell(row=row, column=cell.column).number_format = "@"


def _autosize(sheet, limit: int = 46) -> None:
    for column in range(1, min(sheet.max_column, 60) + 1):
        longest = 0
        for row in range(1, min(sheet.max_row, 400) + 1):
            value = sheet.cell(row=row, column=column).value
            if value is None:
                continue
            longest = max(longest, min(len(str(value)), limit))
        sheet.column_dimensions[get_column_letter(column)].width = max(10, longest + 2)
