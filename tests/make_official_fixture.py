"""產生一份仿照官方「上機編號_說明」對照表版面的測試檔。

版面（欄位偏移）取自真實檔案，資料則是虛構的 —— 測的是解析邏輯，
不需要也不應該把真實研究資料放進版本庫。

真實版面的三個區塊：
    C 欄      採樣時間點（只寫在每組第一列）
    D / E 欄  動物編號，分別在 M / F 標題之下 —— 性別是靠欄位位置表示的
    G–I 欄    組織代號、英文名、中文名
    K–P 欄    批次 × 時間點 的 Run 編號矩陣，其下另有一塊未排定的區塊
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook

FIXTURES = Path(__file__).parent / "fixtures"
OUTPUT = FIXTURES / "official_mapping_table.xlsx"

TIMEPOINTS = ["Pre-dose", "Day 02", "Day 08", "Day 15", "1 hour"]

# (時間點, 公動物編號, 母動物編號)
ANIMALS = [
    ("Pre-dose", "1006", "0006"), ("", "1007", "0007"), ("", "1008", "0008"),
    ("Day 02", "1011", "0011"), ("", "1012", "0012"), ("", "1031", "0013"),
    ("Day 08", "1016", "0016"), ("", "1017", "0017"),
    ("Day 15", "1021", "0021"), ("", "1022", "0022"),
    ("1 hour", "1066", "0066"), ("", "1067", "0067"),
]

ORGANS = [
    ("21", "Heart", "心"), ("19", "Liver", "肝"), ("22", "Spleen", "脾"),
    ("23", "Lung", "肺"), ("24", "Kidney", "腎"), ("04", "Brain", "腦"),
    ("25", "Testis", "睪丸"), ("11", "Ovaries", "卵巢"),
    ("01", "Adrenal gland", "腎上腺"), ("61", "Cervical", "頸椎"),
    ("62", "Lumber", "腰椎"), ("63", "Thoracic", "胸椎"),
    ("64", "Lymph node", "淋巴結"), ("99", "Injection site", "注射處關節"),
    ("03", "Bone marrow", "骨髓pellet"), ("27", "Blood", "血液pellet"),
]

BATCHES = ["99+03+64", "27+23+22", "19+24+21", "04+25+01", "61+62+63"]

# {批次: {時間點: Run 編號}} —— 只有前兩個批次的前三個時間點排定了
SCHEDULED = {
    "99+03+64": {"Pre-dose": 1, "Day 02": 2, "Day 08": 3},
    "27+23+22": {"Pre-dose": 4, "Day 02": 5, "Day 08": 6},
}


def build() -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "工作表1"

    sheet.cell(row=1, column=1, value="BD-TS-TEST-0001")

    # 標題列
    sheet.cell(row=3, column=4, value="動物編號")
    sheet.cell(row=3, column=5, value="動物編號")
    sheet.cell(row=3, column=7, value="組織代號")
    sheet.cell(row=3, column=8, value="臟器名稱")
    sheet.cell(row=3, column=11, value="上機樣品編號: 動物編號_組織代號")

    # 性別是欄位位置，不是欄位值
    sheet.cell(row=4, column=2, value="送樣日期")
    sheet.cell(row=4, column=4, value="M")
    sheet.cell(row=4, column=5, value="F")

    for offset, (timepoint, male, female) in enumerate(ANIMALS):
        row = 5 + offset
        if timepoint:
            sheet.cell(row=row, column=3, value=timepoint)
        sheet.cell(row=row, column=4, value=male)
        # 母動物編號刻意存成數字，模擬 Excel 吃掉前導零的情況
        sheet.cell(row=row, column=5, value=int(female))

    for offset, (code, english, chinese) in enumerate(ORGANS):
        row = 4 + offset
        sheet.cell(row=row, column=7, value=code)
        sheet.cell(row=row, column=8, value=english)
        sheet.cell(row=row, column=9, value=chinese)

    # 主排程矩陣
    for index, timepoint in enumerate(TIMEPOINTS):
        sheet.cell(row=4, column=12 + index, value=timepoint)
    for offset, batch in enumerate(BATCHES):
        row = 5 + offset
        sheet.cell(row=row, column=11, value=batch)
        for index, timepoint in enumerate(TIMEPOINTS):
            run = SCHEDULED.get(batch, {}).get(timepoint)
            if run is not None:
                sheet.cell(row=row, column=12 + index, value=run)

    # 下方另一塊：列了批次但沒填 Run 編號
    sheet.cell(row=11, column=12, value="Day 29(Vehicle)")
    for offset, batch in enumerate(BATCHES):
        sheet.cell(row=12 + offset, column=11, value=batch)

    FIXTURES.mkdir(parents=True, exist_ok=True)
    workbook.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(f"written: {build()}")
