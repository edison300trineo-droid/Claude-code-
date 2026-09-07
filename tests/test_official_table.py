"""官方「上機編號_說明」對照表解析。

這份檔案的性別是靠「動物編號寫在 M 欄還是 F 欄」表示的，不是一個欄位值 ——
用欄名比對永遠抓不到，所以需要專門的解析器，也需要專門的測試。
"""

from pathlib import Path

import pytest
import pandas as pd
from openpyxl import Workbook

from qpcr_biod.official_table import parse_official_table

FIXTURE = Path(__file__).parent / "fixtures" / "official_mapping_table.xlsx"


@pytest.fixture(scope="module")
def table():
    return parse_official_table(FIXTURE)


# --- 動物：性別由欄位位置決定 ---------------------------------------------

def test_sex_comes_from_which_column_the_id_sits_in(table):
    assert table.animals["1006"]["sex"] == "Male"
    assert table.animals["0006"]["sex"] == "Female"


def test_timepoint_carries_down_until_the_next_one(table):
    """時間點只寫在每組第一列，其餘要沿用上一個值。"""
    for animal in ("1006", "1007", "1008", "0007"):
        assert table.animals[animal]["timepoint"] == "Pre-dose"
    for animal in ("1011", "1012", "1031"):
        assert table.animals[animal]["timepoint"] == "Day 02"
    assert table.animals["1066"]["timepoint"] == "1 hour"


def test_leading_zeros_are_restored(table):
    """母動物編號在 Excel 裡存成數字 6，必須還原成 0006。"""
    assert "0006" in table.animals
    assert "6" not in table.animals


def test_out_of_sequence_ids_are_kept(table):
    """1031 出現在 Day 02 的公動物欄，編號不連續也要照收。"""
    assert table.animals["1031"] == {"sex": "Male", "timepoint": "Day 02"}


def test_every_animal_has_both_sex_and_timepoint(table):
    assert len(table.animals) == 24
    assert all(record["sex"] and record["timepoint"]
               for record in table.animals.values())


# --- 臟器代碼 --------------------------------------------------------------

def test_organ_codes_are_read_with_both_names(table):
    assert len(table.organ_codes) == 16
    assert table.organ_codes["27"] == {"en": "Blood", "zh": "血液pellet"}
    assert table.organ_codes["04"]["en"] == "Brain"


def test_two_digit_organ_codes_keep_their_leading_zero(table):
    assert "04" in table.organ_codes
    assert "01" in table.organ_codes


# --- 批次與 Run 排程 -------------------------------------------------------

def test_batches_are_split_into_organ_codes(table):
    assert table.batches["99+03+64"] == ["99", "03", "64"]
    assert len(table.batches) == 5


def test_only_the_scheduled_runs_are_reported(table):
    """官方表只填了 6 格 Run 編號，不能多生出來。"""
    assert len(table.run_schedule) == 6
    by_run = {entry["run"]: entry for entry in table.run_schedule}
    assert by_run[1]["batch"] == "99+03+64"
    assert by_run[1]["timepoint"] == "Pre-dose"
    assert by_run[6]["batch"] == "27+23+22"
    assert by_run[6]["timepoint"] == "Day 08"


def test_a_second_block_without_run_numbers_is_surfaced(table):
    """Day 29(Vehicle) 那塊列了批次卻沒排 Run，不能被靜默略過。"""
    assert "Day 29(Vehicle)" in table.unscheduled_sections
    assert any("Day 29(Vehicle)" in w for w in table.warnings)


# --- 錨點缺失時的行為 -------------------------------------------------------

def _sheet_without(label: str, tmp_path: Path) -> Path:
    from openpyxl import load_workbook
    workbook = load_workbook(FIXTURE)
    sheet = workbook.active
    for row in sheet.iter_rows():
        for cell in row:
            if cell.value == label:
                cell.value = None
    target = tmp_path / "broken.xlsx"
    workbook.save(target)
    return target


def test_missing_animal_anchor_is_reported_not_guessed(tmp_path):
    result = parse_official_table(_sheet_without("動物編號", tmp_path))
    assert not result.animals
    assert any("動物編號" in w for w in result.warnings)
    # 其他區塊仍應正常解析
    assert result.organ_codes and result.batches


def test_missing_sex_markers_stop_animal_parsing(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.cell(row=3, column=4, value="動物編號")
    sheet.cell(row=5, column=4, value="1006")
    target = tmp_path / "no_sex.xlsx"
    workbook.save(target)

    result = parse_official_table(target)
    assert not result.animals
    assert any("M / F" in w for w in result.warnings)


def test_missing_organ_anchor_is_reported(tmp_path):
    result = parse_official_table(_sheet_without("組織代號", tmp_path))
    assert not result.organ_codes
    assert any("組織代號" in w for w in result.warnings)
    assert result.animals


def test_empty_sheet_yields_nothing_without_crashing(tmp_path):
    workbook = Workbook()
    target = tmp_path / "empty.xlsx"
    workbook.save(target)
    result = parse_official_table(target)
    assert result.is_empty
    assert result.warnings


# --- 整合：官方對照表如何影響統整結果 ---------------------------------------

def _use_official_only(study_dir: Path) -> None:
    """把參考資料夾換成只有官方對照表，排除別名比對的干擾。"""
    reference_dir = study_dir / "data/reference"
    for existing in reference_dir.glob("*.xlsx"):
        existing.unlink()
    import shutil
    shutil.copy(FIXTURE, reference_dir / "BD-TS-20260701 上機編號_說明.xlsx")


def test_sex_and_timepoint_reach_the_consolidated_table(study_dir, config):
    from qpcr_biod.pipeline import run_pipeline

    _use_official_only(study_dir)
    table = run_pipeline(config).consolidated

    female = table[table["動物編號"] == "0006"].iloc[0]
    male = table[table["動物編號"] == "1006"].iloc[0]
    assert female["性別(Sex)"] == "Female"
    assert male["性別(Sex)"] == "Male"


def test_timepoints_are_normalised_to_the_reporting_form(study_dir, config):
    """官方表寫 Pre-dose / Day 02，統整表必須收斂成 Predose / Day 2。"""
    from qpcr_biod.pipeline import run_pipeline

    _use_official_only(study_dir)
    table = run_pipeline(config).consolidated
    timepoints = set(table["採樣時間點(Time point)"].dropna())

    assert "Predose" in timepoints
    assert "Pre-dose" not in timepoints, "未正規化的寫法會讓同一時間點被拆成兩個"
    assert "Day 2" in timepoints
    assert "Day 02" not in timepoints


def test_organ_names_come_from_the_official_table(study_dir, config):
    from qpcr_biod.pipeline import run_pipeline

    _use_official_only(study_dir)
    result = run_pipeline(config)
    codes = pd.read_excel(result.output_path, sheet_name="臟器代碼對照表", dtype=object)

    blood = codes[codes["代碼"].astype(str).str.zfill(2) == "27"].iloc[0]
    assert blood["臟器名稱(英文)"] == "Blood"
    assert blood["資料來源"] == "官方上機編號對照表"


def test_batch_composition_comes_from_the_official_table(study_dir, config):
    from qpcr_biod.pipeline import run_pipeline

    _use_official_only(study_dir)
    result = run_pipeline(config)
    assert any("批次組成取自官方上機編號對照表" in w for w in result.warnings)


def test_unscheduled_section_is_reported_to_the_operator(study_dir, config):
    from qpcr_biod.pipeline import run_pipeline

    _use_official_only(study_dir)
    result = run_pipeline(config)
    assert any("Day 29(Vehicle)" in w for w in result.warnings)
